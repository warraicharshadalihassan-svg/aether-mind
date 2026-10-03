import os
from io import BytesIO
from typing import Optional
from fastapi import FastAPI, HTTPException, Header
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
import openai
import anthropic
import google.generativeai as genai
from pptx import Presentation
from pptx.util import Inches, Pt

app = FastAPI(title="Nexora AI Multimodal Router API", version="2.0.0")

# Inicialización de clientes desde variables de entorno
anthropic_client = anthropic.Anthropic(api_key=os.getenv("ANTHROPIC_API_KEY"))
genai.configure(api_key=os.getenv("GEMINI_API_KEY"))
openai_client = openai.OpenAI(api_key=os.getenv("OPENAI_API_KEY"))

# Instrucción unificada de identidad para evitar filtrado de marca
SYSTEM_PROMPT = (
    "Eres Nexora AI, una plataforma de inteligencia artificial unificada de alto rendimiento. "
    "Responde siempre directamente, con tono profesional y eficiente."
)

class ChatRequest(BaseModel):
    message: str
    model: str = "auto"
    stream: bool = False

class PPTRequest(BaseModel):
    topic: str
    slides_count: int = 5

@app.get("/")
def read_root():
    return {"status": "online", "service": "Nexora AI API"}

@app.post("/v1/chat")
async def chat_endpoint(request: ChatRequest, authorization: Optional[str] = Header(None)):
    user_message = request.message
    model_choice = request.model.lower()

    try:
        if model_choice == "claude":
            response = anthropic_client.messages.create(
                model="claude-3-5-sonnet-20241022",
                max_tokens=1024,
                system=SYSTEM_PROMPT,
                messages=[{"role": "user", "content": user_message}]
            )
            return {"response": response.content[0].text, "model_used": "claude-3-5-sonnet"}

        elif model_choice == "gemini":
            model = genai.GenerativeModel("gemini-1.5-flash", system_instruction=SYSTEM_PROMPT)
            response = model.generate_content(user_message)
            return {"response": response.text, "model_used": "gemini-1.5-flash"}

        else:
            response = openai_client.chat.completions.create(
                model="gpt-4o-mini",
                messages=[
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": user_message}
                ]
            )
            return {"response": response.choices[0].message.content, "model_used": "gpt-4o-mini"}

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error en Nexora AI: {str(e)}")

@app.post("/v1/generate-ppt")
async def generate_ppt_endpoint(request: PPTRequest):
    try:
        prs = Presentation()
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]

        title.text = request.topic
        subtitle.text = "Presentación generada automáticamente por Nexora AI"

        for i in range(1, request.slides_count):
            bullet_slide_layout = prs.slide_layouts[1]
            s = prs.slides.add_slide(bullet_slide_layout)
            s.shapes.title.text = f"Punto {i}: {request.topic}"
            tf = s.placeholders[1].text_frame
            tf.text = f"Detalle importante número {i} sobre el tema."

        output = BytesIO()
        prs.save(output)
        output.seek(0)

        headers = {
            "Content-Disposition": f"attachment; filename=Nexora_AI_{request.topic.replace(' ', '_')}.pptx"
        }
        return StreamingResponse(
            output, 
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation", 
            headers=headers
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error generando presentación: {str(e)}")
