import os
from io import BytesIO
from typing import Optional
from fastapi import FastAPI, HTTPException, Header
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
import openai
import anthropic
import google.generativeai as genai
from pptx import Presentation
from pptx.util import Inches, Pt

app = FastAPI(title="Aether Mind Multimodal Router API", version="2.0.0")

# Inicialización de clientes desde variables de entorno
anthropic_client = anthropic.Anthropic(api_key=os.getenv("ANTHROPIC_API_KEY"))
genai.configure(api_key=os.getenv("GEMINI_API_KEY"))
openai_client = openai.OpenAI(api_key=os.getenv("OPENAI_API_KEY"))

# Instrucción unificada de identidad para evitar filtrado de marca
SYSTEM_PROMPT = (
    "Eres Aether Mind, una plataforma de inteligencia artificial unificada de alto rendimiento. "
    "Responde siempre directamente, con tono profesional y eficiente. "
    "Jamás te identifiques como ChatGPT, Claude, Gemini o modelos creados por OpenAI, Anthropic o Google."
)

class ChatRequest(BaseModel):
    user_id: str
    prompt: str

class ChatResponse(BaseModel):
    routed_model: str
    response_type: str  # text, image_url, pptx_file
    content: str

def route_intent(prompt: str) -> str:
    """Enruta la petición al modelo o servicio adecuado según las palabras clave."""
    p = prompt.lower()
    
    # Detección por tipo de tarea
    image_kw = ["imagen", "dibuja", "genera foto", "diseña un logo", "render", "haz una imagen", "dall-e"]
    ppt_kw = ["powerpoint", "presentacion", "diapositivas", "crea un ppt", "pptx"]
    code_kw = ["code", "python", "javascript", "function", "bug", "script", "api", "react", "html", "css", "sql"]
    doc_kw = ["pdf", "summarize", "documento", "analizar texto", "resumen", "contexto", "archivo"]
    
    if any(k in p for k in image_kw):
        return "aether-vision-generator"
    elif any(k in p for k in ppt_kw):
        return "aether-slides-generator"
    elif any(k in p for k in code_kw):
        return "claude-3-5-sonnet-latest"
    elif any(k in p for k in doc_kw):
        return "gemini-1.5-pro-latest"
    else:
        return "gpt-4o"

def build_presentation(topic: str) -> BytesIO:
    """Genera una presentación PowerPoint básica a partir del tema solicitado."""
    prs = Presentation()
    
    # Diapositiva 1: Portada
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = topic.capitalize()
    slide.placeholders[1].text = "Presentación creada automáticamente por Aether Mind"
    
    # Diapositiva 2: Contenido
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Puntos Clave"
    tf = slide.placeholders[1].text_frame
    tf.text = f"Análisis inicial sobre: {topic}"
    p = tf.add_paragraph()
    p.text = "Generación de contenido optimizada con arquitectura Aether Mind."
    
    buffer = BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer

@app.post("/api/v1/chat", response_model=ChatResponse)
async def process_chat(request: ChatRequest, authorization: Optional[str] = Header(None)):
    if not authorization:
        raise HTTPException(status_code=401, detail="Token de autorización no proporcionado.")
    
    selected_model = route_intent(request.prompt)

    try:
        # 1. Generación de Imágenes (DALL-E 3)
        if selected_model == "aether-vision-generator":
            img_res = openai_client.images.generate(
                model="dall-e-3",
                prompt=f"{request.prompt}. High quality, professional style.",
                n=1,
                size="1024x1024"
            )
            return ChatResponse(
                routed_model=selected_model,
                response_type="image_url",
                content=img_res.data[0].url
            )

        # 2. Generación de Presentaciones (PowerPoint)
        elif selected_model == "aether-slides-generator":
            return ChatResponse(
                routed_model=selected_model,
                response_type="pptx_file",
                content="Llama al endpoint GET /api/v1/download-ppt con tu tema para descargar el archivo .pptx"
            )

        # 3. Código y Desarrollo (Claude)
        elif selected_model == "claude-3-5-sonnet-latest":
            response = anthropic_client.messages.create(
                model="claude-3-5-sonnet-latest",
                max_tokens=4096,
                system=SYSTEM_PROMPT,
                messages=[{"role": "user", "content": request.prompt}]
            )
            return ChatResponse(
                routed_model=selected_model,
                response_type="text",
                content=response.content[0].text
            )

        # 4. Análisis de Documentos y Contextos Largos (Gemini Pro)
        elif selected_model == "gemini-1.5-pro-latest":
            model = genai.GenerativeModel('gemini-1.5-pro-latest', system_instruction=SYSTEM_PROMPT)
            response = model.generate_content(request.prompt)
            return ChatResponse(
                routed_model=selected_model,
                response_type="text",
                content=response.text
            )

        # 5. Modelo General / Razonamiento (GPT-4o)
        else:
            response = openai_client.chat.completions.create(
                model="gpt-4o",
                messages=[
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": request.prompt}
                ]
            )
            return ChatResponse(
                routed_model=selected_model,
                response_type="text",
                content=response.choices[0].message.content
            )

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error en Aether Mind Engine: {str(e)}")

@app.get("/api/v1/download-ppt")
async def download_ppt(topic: str):
    """Endpoint directo para descargar un archivo .pptx generado."""
    ppt_buffer = build_presentation(topic)
    headers = {'Content-Disposition': f'attachment; filename="{topic}_aether_mind.pptx"'}
    return StreamingResponse(
        ppt_buffer,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers=headers
    )as e:
        raise HTTPException(status_code=500, detail=f"Error en la API de IA: {str(e)}")
